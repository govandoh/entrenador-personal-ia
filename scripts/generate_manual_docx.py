# -*- coding: utf-8 -*-
"""
Genera docs/manual-usuario.docx a partir del contenido del manual de usuario.
Formato: manual de usuario académico con estilos UMG.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT_PATH = os.path.join(os.path.dirname(__file__), '..', 'docs', 'manual-usuario.docx')

# ──────────────────────────────────────────────
# Helpers de formato (compartidos con generate_docx.py)
# ──────────────────────────────────────────────

def set_font(run, name='Times New Roman', size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)

def set_paragraph_spacing(para, before=0, after=6, line_rule=WD_LINE_SPACING.ONE_POINT_FIVE):
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after  = Pt(after)
    pf.line_spacing_rule = line_rule

def add_bottom_border(para):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),   'single')
    bottom.set(qn('w:sz'),    '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '003366')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_heading(doc, text, level=1):
    para = doc.add_paragraph()
    run  = para.add_run(text)
    if level == 1:
        set_font(run, size=14, bold=True, color=(0, 51, 102))
        para.paragraph_format.space_before = Pt(18)
        para.paragraph_format.space_after  = Pt(6)
        add_bottom_border(para)
    elif level == 2:
        set_font(run, size=13, bold=True, color=(0, 70, 127))
        para.paragraph_format.space_before = Pt(14)
        para.paragraph_format.space_after  = Pt(4)
    elif level == 3:
        set_font(run, size=12, bold=True, italic=True, color=(40, 40, 40))
        para.paragraph_format.space_before = Pt(10)
        para.paragraph_format.space_after  = Pt(3)
    return para

def add_body(doc, text, indent=False):
    para = doc.add_paragraph()
    run  = para.add_run(text)
    set_font(run, size=11)
    set_paragraph_spacing(para, before=0, after=6)
    para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        para.paragraph_format.left_indent = Cm(0.5)
    return para

def add_body_mixed(doc, segments, indent=False):
    """Párrafo con mezcla de bold/italic inline. segments: [(texto, bold, italic)]"""
    para = doc.add_paragraph()
    for text, bold, italic in segments:
        run = para.add_run(text)
        set_font(run, size=11, bold=bold, italic=italic)
    set_paragraph_spacing(para, before=0, after=6)
    para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        para.paragraph_format.left_indent = Cm(0.5)
    return para

def add_bullet(doc, text, level=0, bold_prefix=None):
    """Bullet con opcional prefijo en negrita (bold_prefix: texto antes del ':')"""
    para = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r_pre = para.add_run(f'{bold_prefix}: ')
        set_font(r_pre, size=11, bold=True)
    run = para.add_run(text)
    set_font(run, size=11)
    para.paragraph_format.left_indent  = Cm(1.0 + level * 0.5)
    para.paragraph_format.space_after  = Pt(3)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return para

def add_numbered(doc, text, num, bold_prefix=None):
    para = doc.add_paragraph()
    run_num = para.add_run(f'{num}. ')
    set_font(run_num, size=11, bold=True)
    if bold_prefix:
        r_pre = para.add_run(f'{bold_prefix}: ')
        set_font(r_pre, size=11, bold=True)
    run_text = para.add_run(text)
    set_font(run_text, size=11)
    para.paragraph_format.left_indent  = Cm(0.5)
    para.paragraph_format.space_after  = Pt(4)
    para.paragraph_format.alignment    = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return para

def add_note(doc, text, kind='note'):
    """
    Caja de nota/advertencia/alerta.
    kind: 'note' (azul), 'warning' (amarillo), 'alert' (rojo)
    """
    colors  = {'note': 'D6EAF8', 'warning': 'FEF9E7', 'alert': 'FDEDEC'}
    borders = {'note': '2980B9', 'warning': 'F39C12', 'alert': 'E74C3C'}
    labels  = {'note': 'Nota', 'warning': 'Consejo', 'alert': 'Importante'}
    fill   = colors.get(kind,  'D6EAF8')
    border = borders.get(kind, '2980B9')
    label  = labels.get(kind,  'Nota')

    para = doc.add_paragraph()
    r_label = para.add_run(f'{label}: ')
    set_font(r_label, size=11, bold=True)
    r_text = para.add_run(text)
    set_font(r_text, size=11)
    para.paragraph_format.left_indent  = Cm(0.5)
    para.paragraph_format.right_indent = Cm(0.5)
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after  = Pt(4)
    para.paragraph_format.alignment    = WD_ALIGN_PARAGRAPH.JUSTIFY

    pPr = para._p.get_or_add_pPr()
    # Fondo
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  fill)
    pPr.append(shd)
    # Borde izquierdo
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'),   'single')
    left.set(qn('w:sz'),    '12')
    left.set(qn('w:space'), '4')
    left.set(qn('w:color'), border)
    pBdr.append(left)
    pPr.append(pBdr)
    return para

def add_screenshot_placeholder(doc, description):
    """Bloque gris con borde punteado para indicar dónde va una captura de pantalla."""
    para = doc.add_paragraph()
    run  = para.add_run(f'[ CAPTURA DE PANTALLA: {description} ]')
    run.font.name    = 'Calibri'
    run.font.size    = Pt(10)
    run.font.italic  = True
    run.font.color.rgb = RGBColor(100, 100, 100)
    para.paragraph_format.alignment    = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after  = Pt(6)
    para.paragraph_format.left_indent  = Cm(1.0)
    para.paragraph_format.right_indent = Cm(1.0)

    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  'EFEFEF')
    pPr.append(shd)
    pBdr = OxmlElement('w:pBdr')
    for side in ('top', 'bottom', 'left', 'right'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'),   'dashed')
        el.set(qn('w:sz'),    '6')
        el.set(qn('w:space'), '4')
        el.set(qn('w:color'), 'AAAAAA')
        pBdr.append(el)
    pPr.append(pBdr)
    return para

def make_table(doc, headers, rows, col_widths=None):
    """Tabla genérica con encabezado azul oscuro."""
    n_cols = len(headers)
    n_rows = len(rows) + 1
    t = doc.add_table(rows=n_rows, cols=n_cols)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Encabezado
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        r = cell.paragraphs[0].add_run(h)
        set_font(r, size=10, bold=True, color=(255, 255, 255))
        cell.paragraphs[0].paragraph_format.space_after  = Pt(2)
        cell.paragraphs[0].paragraph_format.space_before = Pt(2)
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'),   'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'),  '003366')
        tcPr.append(shd)

    # Filas de datos
    for i, row_data in enumerate(rows, 1):
        # Alternar fondo en filas pares para legibilidad
        fill_color = 'EAF2FB' if i % 2 == 0 else 'FFFFFF'
        for j, val in enumerate(row_data):
            cell = t.rows[i].cells[j]
            p    = cell.paragraphs[0]
            # Soporte para bold inline: texto entre ** se pone en negrita
            if '**' in val:
                parts = val.split('**')
                for k, part in enumerate(parts):
                    if part:
                        rr = p.add_run(part)
                        set_font(rr, size=10, bold=(k % 2 == 1))
            else:
                rr = p.add_run(val)
                set_font(rr, size=10)
            p.paragraph_format.space_after  = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            if fill_color != 'FFFFFF':
                tc   = cell._tc
                tcPr = tc.get_or_add_tcPr()
                shd  = OxmlElement('w:shd')
                shd.set(qn('w:val'),   'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'),  fill_color)
                tcPr.append(shd)

    # Ajuste de ancho de columnas si se especifica
    if col_widths:
        for i, row in enumerate(t.rows):
            for j, w in enumerate(col_widths):
                if j < len(row.cells):
                    row.cells[j].width = Cm(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def add_page_numbers(doc):
    for section in doc.sections:
        footer = section.footer
        para   = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        para.clear()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        fldChar1  = OxmlElement('w:fldChar');  fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText'); instrText.text = ' PAGE '
        fldChar2  = OxmlElement('w:fldChar');  fldChar2.set(qn('w:fldCharType'), 'end')
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)

def set_page_margins(doc):
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(2.5)

def hr(doc):
    """Línea horizontal separadora."""
    para = doc.add_paragraph()
    pPr  = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    top  = OxmlElement('w:top')
    top.set(qn('w:val'),   'single')
    top.set(qn('w:sz'),    '4')
    top.set(qn('w:space'), '1')
    top.set(qn('w:color'), 'CCCCCC')
    pBdr.append(top)
    pPr.append(pBdr)
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after  = Pt(6)

# ──────────────────────────────────────────────
# Construcción del documento
# ──────────────────────────────────────────────

doc = Document()
set_page_margins(doc)
add_page_numbers(doc)

# ══════════════════════════════════════════════
# PORTADA
# ══════════════════════════════════════════════
doc.add_paragraph()

for txt, sz, bold, color in [
    ('UNIVERSIDAD MARIANO GÁLVEZ DE GUATEMALA', 14, True,  (0, 51, 102)),
    ('Facultad de Ingeniería en Sistemas de Información',  12, False, None),
    ('Inteligencia Artificial — IA26',                     12, False, None),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt)
    set_font(r, size=sz, bold=bold, color=color)

doc.add_paragraph(); doc.add_paragraph(); doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('MANUAL DE USUARIO')
set_font(r, size=20, bold=True, color=(0, 51, 102))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Entrenador IA — Entrenador Personal con Estimación de Poses en Tiempo Real')
set_font(r, size=13, italic=True)

doc.add_paragraph(); doc.add_paragraph(); doc.add_paragraph()

for label, value in [
    ('Versión de la aplicación:', '1.0'),
    ('Fecha:', 'Mayo 2026'),
    ('URL de la aplicación:', '(enlace de Vercel)'),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(f'{label} '); set_font(r1, size=11, bold=True)
    r2 = p.add_run(value);       set_font(r2, size=11)
    p.paragraph_format.space_after = Pt(3)

doc.add_page_break()

# ══════════════════════════════════════════════
# ÍNDICE
# ══════════════════════════════════════════════
add_heading(doc, 'Tabla de contenido', level=1)

toc = [
    ('1.',    'Introducción'),
    ('2.',    'Requisitos del dispositivo'),
    ('3.',    'Acceso e instalación'),
    ('   3.1','Abrir desde el navegador'),
    ('   3.2','Instalar como aplicación (PWA) en Android'),
    ('   3.3','Instalar como aplicación (PWA) en iOS'),
    ('4.',    'Primera vez: flujo de bienvenida'),
    ('   4.1','Pantalla de carga (Splash)'),
    ('   4.2','Pantalla "Cómo funciona"'),
    ('   4.3','Pantalla de permisos'),
    ('   4.4','Pantalla de inicio'),
    ('5.',    'Interfaz principal'),
    ('   5.1','Vista de cámara y esqueleto'),
    ('   5.2','Barra de retroalimentación'),
    ('   5.3','Selector de ejercicio'),
    ('   5.4','Botón de cambio de cámara'),
    ('6.',    'Ejercicios disponibles'),
    ('   6.1','Sentadilla'),
    ('   6.2','Curl de Bíceps'),
    ('   6.3','Press de Hombro'),
    ('7.',    'Sistema de retroalimentación'),
    ('   7.1','Retroalimentación visual'),
    ('   7.2','Retroalimentación de voz'),
    ('8.',    'Cambiar de ejercicio'),
    ('9.',    'Cambiar de cámara durante el entrenamiento'),
    ('10.',   'Solución de problemas'),
    ('11.',   'Preguntas frecuentes'),
]
for num, title in toc:
    p = doc.add_paragraph()
    main = not num.startswith(' ')
    r1 = p.add_run(f'{num.strip()}  '); set_font(r1, size=11, bold=main)
    r2 = p.add_run(title);              set_font(r2, size=11, bold=main)
    p.paragraph_format.left_indent  = Cm(0 if main else 0.8)
    p.paragraph_format.space_after  = Pt(2)

doc.add_page_break()

# ══════════════════════════════════════════════
# 1. INTRODUCCIÓN
# ══════════════════════════════════════════════
add_heading(doc, '1. Introducción', level=1)
add_body(doc,
    'Entrenador IA es una aplicación web progresiva (PWA) gratuita que utiliza la cámara de tu celular para '
    'analizar tu postura corporal en tiempo real. Mediante el modelo de inteligencia artificial MediaPipe Pose '
    'de Google, la aplicación detecta 33 puntos clave de tu cuerpo en cada fotograma y calcula los ángulos '
    'de tus articulaciones para evaluar la calidad de tus ejercicios.'
)
add_body(doc, 'La aplicación:')
for item in [
    'Cuenta repeticiones automáticamente.',
    'Emite retroalimentación visual en colores (verde, amarillo, rojo) según tu técnica.',
    'Habla en voz alta para guiarte sin que tengas que mirar la pantalla.',
    'Funciona completamente en tu dispositivo: el video de tu cámara nunca sale de tu celular.',
    'No requiere cuenta de usuario, suscripción ni conexión permanente a internet.',
]:
    add_bullet(doc, item)

add_note(doc, '¿Qué necesito? Un teléfono inteligente con cámara, navegador moderno y buena iluminación.', kind='note')

# ══════════════════════════════════════════════
# 2. REQUISITOS
# ══════════════════════════════════════════════
add_heading(doc, '2. Requisitos del dispositivo', level=1)
make_table(doc,
    headers=['Requisito', 'Detalle'],
    rows=[
        ['**Sistema operativo**', 'Android 9 o superior / iOS 14.5 o superior'],
        ['**Navegador**',         'Chrome 90+ (Android)  ·  Safari 15+ (iOS)  ·  Firefox 90+'],
        ['**Cámara**',            'Cámara trasera o frontal funcional'],
        ['**Conexión**',          'Necesaria solo la primera vez para descargar el modelo de IA (~8 MB)'],
        ['**Iluminación**',       'Ambiente bien iluminado (luz natural o artificial uniforme)'],
        ['**Espacio**',           'Área libre donde el cuerpo completo sea visible desde la cintura hasta los pies'],
    ],
    col_widths=[4.5, 11.5],
)
add_note(doc,
    'La aplicación no es compatible con navegadores privados o sin cookies en iOS, ya que requiere '
    'almacenamiento local para recordar tu configuración.',
    kind='warning',
)

# ══════════════════════════════════════════════
# 3. ACCESO E INSTALACIÓN
# ══════════════════════════════════════════════
add_heading(doc, '3. Acceso e instalación', level=1)

add_heading(doc, '3.1 Abrir desde el navegador', level=2)
for i, txt in enumerate([
    'Abre el navegador de tu celular (Chrome en Android, Safari en iOS).',
    'Escribe o pega la URL de la aplicación: (enlace de Vercel).',
    'La aplicación carga directamente, sin necesidad de instalación previa.',
], 1):
    add_numbered(doc, txt, i)

hr(doc)

add_heading(doc, '3.2 Instalar como aplicación (PWA) en Android', level=2)
add_body(doc,
    'Instalar la app te permite abrirla desde tu pantalla de inicio como cualquier otra aplicación, '
    'en pantalla completa y sin barra del navegador.'
)
add_screenshot_placeholder(doc, 'menú "Instalar aplicación" o "Añadir a pantalla de inicio" en Chrome Android')
for i, txt in enumerate([
    'Abre la app en Chrome en tu teléfono Android.',
    'Toca el ícono de tres puntos (⋮) en la esquina superior derecha del navegador.',
    'Selecciona "Instalar aplicación" o "Añadir a pantalla de inicio".',
    'En el diálogo que aparece, toca "Instalar".',
    'El ícono de Entrenador IA aparecerá en tu pantalla de inicio.',
], 1):
    add_numbered(doc, txt, i)

hr(doc)

add_heading(doc, '3.3 Instalar como aplicación (PWA) en iOS', level=2)
add_screenshot_placeholder(doc, 'menú Compartir de Safari con opción "Agregar a pantalla de inicio" en iOS')
for i, txt in enumerate([
    'Abre la app en Safari en tu iPhone o iPad.',
    'Toca el botón Compartir (el ícono de caja con flecha hacia arriba, en la barra inferior del navegador).',
    'Desplázate hacia abajo en el menú y toca "Agregar a pantalla de inicio".',
    'Cambia el nombre si lo deseas y toca "Agregar" en la esquina superior derecha.',
    'El ícono de Entrenador IA aparecerá en tu pantalla de inicio.',
], 1):
    add_numbered(doc, txt, i)
add_note(doc,
    'Importante para iOS: Siempre abre la app desde Safari la primera vez. '
    'Una vez instalada, ábrela desde su ícono en la pantalla de inicio.',
    kind='alert',
)

# ══════════════════════════════════════════════
# 4. FLUJO DE BIENVENIDA
# ══════════════════════════════════════════════
add_heading(doc, '4. Primera vez: flujo de bienvenida', level=1)
add_body(doc,
    'La primera vez que abres la aplicación, un flujo de cuatro pantallas te guía para configurarla '
    'correctamente. Este flujo solo aparece una vez; en las siguientes aperturas irás directamente a la cámara.'
)

add_heading(doc, '4.1 Pantalla de carga (Splash)', level=2)
add_screenshot_placeholder(doc, 'pantalla de splash con logo "Entrenador IA" y loader de puntos animados')
add_body(doc,
    'Al abrir la app, aparece una pantalla con el logotipo de Entrenador IA y una animación de carga. '
    'Esta pantalla avanza automáticamente después de aproximadamente 3 segundos; '
    'no requiere ninguna acción de tu parte.'
)

add_heading(doc, '4.2 Pantalla "Cómo funciona"', level=2)
add_screenshot_placeholder(doc, 'pantalla "Tu cuerpo, analizado en tiempo real" con los tres pasos ilustrados')
add_body(doc, 'Esta pantalla explica el funcionamiento básico de la aplicación en tres pasos:')
make_table(doc,
    headers=['Paso', 'Descripción'],
    rows=[
        ['**1. Apunta tu cámara**',   'Posiciona tu cuerpo completo en el encuadre y empieza el ejercicio.'],
        ['**2. Detectamos tu pose**',  'La IA analiza ángulos articulares en cada fotograma, sin enviar datos a ningún servidor.'],
        ['**3. Mejora tu técnica**',   'Feedback visual en verde, amarillo o rojo según tu forma. Conteo automático de repeticiones.'],
    ],
    col_widths=[5, 11],
)
add_body(doc, 'Toca "Continuar" para avanzar a la siguiente pantalla.')

add_heading(doc, '4.3 Pantalla de permisos', level=2)
add_screenshot_placeholder(doc, 'pantalla de permisos con tarjetas de Cámara (requerida) y Notificaciones (opcional)')
add_body(doc, 'La aplicación solicita dos permisos:')
make_table(doc,
    headers=['Permiso', 'Tipo', 'Para qué se usa'],
    rows=[
        ['**Cámara**',          'Requerido', 'Detectar tu cuerpo y calcular ángulos articulares en tiempo real.'],
        ['**Notificaciones**',  'Opcional',  'Recordatorios de entrenamiento y resumen de sesión al finalizar.'],
    ],
    col_widths=[4, 3, 9],
)
add_body(doc, 'Al tocar "Dar permisos y continuar", el navegador mostrará el diálogo nativo de permisos de cámara.')
add_screenshot_placeholder(doc, 'diálogo nativo del navegador solicitando permiso de cámara')
add_bullet(doc, 'Toca "Permitir" para habilitar la cámara.')
add_bullet(doc,
    'Si tocas "Denegar", la aplicación no podrá detectar tu pose. Podés cambiar este permiso más tarde '
    'desde la configuración de tu navegador (ver sección 10. Solución de problemas).'
)
add_note(doc,
    'Privacidad: Todo el procesamiento ocurre en tu dispositivo. '
    'El video de tu cámara nunca se envía a ningún servidor.',
    kind='note',
)

add_heading(doc, '4.4 Pantalla de inicio', level=2)
add_screenshot_placeholder(doc, 'pantalla "¡Empecemos a entrenar!" con selector de cámara trasera/frontal activo')
add_body(doc, 'Esta pantalla te permite elegir la cámara con la que vas a entrenar:')
make_table(doc,
    headers=['Opción', 'Cuándo usarla'],
    rows=[
        ['**Trasera** (recomendada)',
         'Cuando apoyás el celular frente a ti: en una silla, repisa o soporte. '
         'Ofrece mayor campo visual y mejor calidad de imagen.'],
        ['**Frontal**',
         'Cuando sostenés el celular en la mano o lo apoyás con la pantalla hacia ti. '
         'Útil para ejercicios de brazos desde cerca.'],
    ],
    col_widths=[5, 11],
)
add_note(doc, 'Podés cambiar de cámara en cualquier momento mientras entrenás, sin necesidad de reiniciar la app.', kind='note')
add_body(doc, 'Toca "Comenzar a entrenar" para entrar a la pantalla principal.')
add_note(doc,
    'Nota técnica: Este botón también desbloquea la función de voz en iOS Safari, '
    'que requiere una acción previa del usuario para emitir audio.',
    kind='warning',
)

# ══════════════════════════════════════════════
# 5. INTERFAZ PRINCIPAL
# ══════════════════════════════════════════════
add_heading(doc, '5. Interfaz principal', level=1)
add_body(doc, 'Una vez completado el onboarding, verás la pantalla principal de entrenamiento.')
add_screenshot_placeholder(doc,
    'pantalla principal anotada con: (A) video cámara, (B) esqueleto superpuesto, '
    '(C) barra de feedback, (D) chips de ejercicio, (E) botón de cámara'
)
add_body(doc, 'La interfaz tiene cinco elementos:')
make_table(doc,
    headers=['Elemento', 'Descripción'],
    rows=[
        ['**(A) Video de cámara**',            'Muestra en tiempo real la imagen de tu cámara.'],
        ['**(B) Esqueleto superpuesto**',       'Líneas verdes y puntos que muestran los 33 puntos corporales detectados.'],
        ['**(C) Barra de retroalimentación**',  'Mensaje de calidad de ejecución + contador de repeticiones, fija en la parte inferior.'],
        ['**(D) Chips de ejercicio**',          'Selector deslizable para cambiar entre los tres ejercicios disponibles.'],
        ['**(E) Botón de cámara**',             'Alterna entre cámara trasera y frontal.'],
    ],
    col_widths=[5.5, 10.5],
)

add_heading(doc, '5.1 Vista de cámara y esqueleto', level=2)
add_screenshot_placeholder(doc, 'cámara activa con esqueleto de MediaPipe (líneas verdes y puntos) visible sobre el cuerpo del usuario')
add_body(doc, 'Cuando la aplicación está activa y detecta tu cuerpo:')
add_bullet(doc, 'Aparecen líneas verdes que conectan las articulaciones detectadas (esqueleto).')
add_bullet(doc, 'Aparecen puntos rojos/amarillos en cada una de las 33 articulaciones detectadas.')
add_body(doc,
    'Si la detección no es correcta (iluminación deficiente, cuerpo fuera de encuadre), el esqueleto puede '
    'aparecer distorsionado o desaparecer. En ese caso, la barra inferior mostrará: '
    '"Asegúrate de que tu cuerpo completo sea visible".'
)
add_note(doc,
    'Mantén al menos 1.5 metros de distancia entre la cámara y tu cuerpo para que las articulaciones '
    'principales sean visibles.',
    kind='warning',
)

add_heading(doc, '5.2 Barra de retroalimentación', level=2)
add_screenshot_placeholder(doc, 'barra de retroalimentación en estado verde con mensaje "¡Excelente profundidad!" y contador "3 REPS"')
add_body(doc, 'La barra inferior siempre visible contiene:')
add_bullet(doc, 'Mensaje de texto: instrucción o evaluación de la ejecución actual.')
add_bullet(doc, 'Contador de repeticiones: número grande que se anima con cada nueva repetición completada.')
add_bullet(doc, 'Borde izquierdo colorido: indica el nivel de calidad de la ejecución (ver sección 7).')

add_heading(doc, '5.3 Selector de ejercicio', level=2)
add_screenshot_placeholder(doc, 'chips de ejercicio en la barra inferior con "Sentadillas" activo (resaltado)')
add_body(doc, 'En la parte inferior de la pantalla hay tres chips deslizables, uno por ejercicio:')
make_table(doc,
    headers=['Chip', 'Ejercicio'],
    rows=[
        ['Ícono de persona en cuclillas', '**Sentadillas**'],
        ['Ícono de barra con discos',     '**Curl de Bíceps**'],
        ['Ícono de persona con brazos levantados', '**Press de Hombro**'],
    ],
    col_widths=[7, 9],
)
add_bullet(doc, 'El chip del ejercicio activo aparece resaltado (fondo blanco, texto oscuro).')
add_bullet(doc, 'Los chips inactivos aparecen con fondo semitransparente.')
add_bullet(doc, 'Si los chips no caben en la pantalla, podés deslizar horizontalmente para ver todos.')

add_heading(doc, '5.4 Botón de cambio de cámara', level=2)
add_screenshot_placeholder(doc, 'botón de cámara circular en esquina de la barra de controles')
add_body(doc,
    'El botón circular con el ícono de cámara con flechas alterna entre la cámara trasera y la cámara frontal. '
    'Durante el cambio, aparece brevemente el mensaje "Cambiando cámara..." mientras el hardware del dispositivo '
    'libera la cámara anterior y activa la nueva.'
)

# ══════════════════════════════════════════════
# 6. EJERCICIOS
# ══════════════════════════════════════════════
add_heading(doc, '6. Ejercicios disponibles', level=1)

# ── 6.1 SENTADILLA ──────────────────────────
add_heading(doc, '6.1 Sentadilla', level=2)
add_screenshot_placeholder(doc, 'chip "Sentadillas" activo y retroalimentación verde "¡Excelente profundidad!"')

add_heading(doc, 'Configuración de la cámara', level=3)
add_bullet(doc,
    'Posición recomendada: Cámara trasera colocada a la altura de la cadera, apuntando '
    'lateralmente o en ángulo de 45°. La vista de frente también funciona, pero la lateral '
    'permite mejor evaluación de la profundidad.',
)
add_bullet(doc,
    'Distancia: Al menos 1.5–2 metros entre la cámara y tu cuerpo para que caderas, '
    'rodillas y tobillos sean visibles simultáneamente.',
)
add_bullet(doc, 'Articulaciones que monitorea: Caderas, rodillas y tobillos de ambas piernas.')

add_heading(doc, 'Cómo ejecutar la sentadilla', level=3)
add_screenshot_placeholder(doc, 'posición inicial de pie con barra inferior idle "Listo — baja para la sentadilla"')
for i, txt in enumerate([
    'Colócate frente a la cámara con el cuerpo completo visible.',
    'Párate erguido. La barra inferior mostrará: "Listo — baja para la sentadilla" (color blanco).',
    'Baja despacio doblando las rodillas, como si fueras a sentarte.',
    'Desciende hasta que el ángulo de tus rodillas sea menor a 100°.',
], 1):
    add_numbered(doc, txt, i)
add_screenshot_placeholder(doc, 'posición de bajada parcial con barra amarilla "Baja un poco más"')
add_numbered(doc, 'Si la barra muestra "Baja un poco más" (amarillo), sigue bajando para lograr mayor profundidad.', 5)
add_screenshot_placeholder(doc, 'posición de profundidad óptima (muslos paralelos al suelo) con barra verde "¡Excelente profundidad!"')
for i, txt in enumerate([
    'Cuando alcanzás muslos paralelos al suelo (ángulo ≤ 90°), la barra cambia a verde: "¡Excelente profundidad!".',
    'Sube hasta la posición inicial. La repetición se contabiliza en ese momento.',
    'La voz anuncia el número de la repetición completada.',
], 6):
    add_numbered(doc, txt, i)

add_heading(doc, 'Tabla de retroalimentación — Sentadilla', level=3)
make_table(doc,
    headers=['Estado', 'Color', 'Mensaje', 'Significado'],
    rows=[
        ['De pie, listo',                    'Blanco',   '"Listo — baja para la sentadilla"',                     'Posición inicial correcta'],
        ['Bajando, profundidad insuficiente', 'Amarillo', '"Baja un poco más"',                                    'Ángulo de rodilla entre 90° y 100°'],
        ['Bajando, profundidad óptima',       'Verde',    '"¡Excelente profundidad!"',                             'Ángulo de rodilla ≤ 90°'],
        ['Cuerpo fuera de encuadre',          'Blanco',   '"Asegúrate de que tu cuerpo completo sea visible"',     'Articulaciones no detectadas'],
    ],
    col_widths=[4, 2.5, 5.5, 4],
)

add_heading(doc, 'Consejos de posición', level=3)
add_bullet(doc, 'Mantén la espalda recta y el pecho hacia adelante.')
add_bullet(doc, 'Las rodillas deben seguir la dirección de los pies, sin colapsar hacia adentro.')
add_bullet(doc, 'Pies separados al ancho de los hombros o un poco más.')

hr(doc)

# ── 6.2 CURL DE BÍCEPS ──────────────────────
add_heading(doc, '6.2 Curl de Bíceps', level=2)
add_screenshot_placeholder(doc, 'chip "Curl de Bíceps" activo con barra verde "¡Contracción completa!"')

add_heading(doc, 'Configuración de la cámara', level=3)
add_body(doc, 'La aplicación detecta automáticamente la vista desde la que estás entrenando:')
make_table(doc,
    headers=['Vista', 'Configuración', 'Cuándo usarla'],
    rows=[
        ['**Frontal**',  'Cámara apuntando de frente a tu cuerpo',    'Curl con barra o mancuernas bilaterales'],
        ['**Lateral**',  'Cámara apuntando al costado de tu cuerpo',  'Curl con mancuerna unilateral; la app detecta automáticamente qué brazo usar'],
    ],
    col_widths=[3, 5.5, 7.5],
)
add_bullet(doc, 'Distancia: Al menos 1 metro para que hombros, codos y muñecas sean visibles.')
add_bullet(doc, 'Articulaciones que monitorea: Hombros, codos y muñecas de ambos brazos.')

add_heading(doc, 'Cómo ejecutar el curl de bíceps', level=3)
add_screenshot_placeholder(doc, 'posición inicial con brazos extendidos y barra idle "Listo — sube el peso"')
for i, txt in enumerate([
    'Párate erguido con los brazos extendidos hacia abajo, sosteniendo el peso (barra, mancuernas o banda).',
    'La barra inferior mostrará: "Listo — sube el peso".',
    'Dobla los codos levantando el peso hacia tus hombros.',
], 1):
    add_numbered(doc, txt, i)
add_screenshot_placeholder(doc, 'posición de máxima contracción con barra verde "¡Contracción completa!"')
for i, txt in enumerate([
    'Cuando el ángulo del codo llega a menos de 50°, la barra muestra "¡Contracción completa!" (verde).',
    'Si la barra muestra "Sube un poco más" (amarillo), el ángulo del codo está entre 50° y 60°.',
    'Baja el peso controladamente hasta extender completamente el brazo. La repetición se registra en ese momento.',
], 4):
    add_numbered(doc, txt, i)

add_heading(doc, 'Modos de uso', level=3)

add_body_mixed(doc, [('Mancuernas bilaterales (ambos brazos simultáneamente):', True, False)])
add_screenshot_placeholder(doc, 'curl bilateral con ambos brazos visibles de frente')
add_bullet(doc, 'Ambos brazos se procesan en paralelo.')
add_bullet(doc, 'La repetición se contabiliza cuando cualquiera de los dos brazos completa el ciclo.')
add_bullet(doc, 'Un mecanismo de cooldown de ~250 ms evita que el segundo brazo genere una repetición doble.')

add_body_mixed(doc, [('Mancuernas alternas (un brazo por vez):', True, False)])
add_screenshot_placeholder(doc, 'curl alterno con un brazo en contracción y el otro extendido')
add_bullet(doc, 'La aplicación detecta y monitorea ambos brazos por separado.')
add_bullet(doc, 'Cada brazo puede completar su repetición de forma independiente con más de 250 ms de diferencia.')

add_body_mixed(doc, [('Vista lateral (un solo brazo):', True, False)])
add_screenshot_placeholder(doc, 'curl lateral con un solo brazo visible desde el costado')
add_bullet(doc,
    'Si la diferencia de visibilidad entre ambos brazos es significativa, la app detecta '
    'automáticamente que estás en vista lateral y procesa únicamente el brazo más visible.'
)

add_heading(doc, 'Tabla de retroalimentación — Curl de Bíceps', level=3)
make_table(doc,
    headers=['Estado', 'Color', 'Mensaje', 'Significado'],
    rows=[
        ['Brazo extendido, listo',       'Blanco',   '"Listo — sube el peso"',          'Posición inicial correcta'],
        ['Subiendo, contracción incompleta', 'Amarillo', '"Sube un poco más"',           'Ángulo del codo entre 50° y 60°'],
        ['Contracción completa',         'Verde',    '"¡Contracción completa!"',         'Ángulo del codo ≤ 50°'],
        ['Brazo fuera de encuadre',      'Blanco',   '"Asegúrate de que tu brazo sea visible"', 'Articulaciones no detectadas'],
    ],
    col_widths=[4.5, 2.5, 5, 4],
)

add_heading(doc, 'Consejos de posición', level=3)
add_bullet(doc, 'Mantén los codos pegados a los costados del torso durante todo el movimiento.')
add_bullet(doc, 'Evita balancear el cuerpo hacia atrás para ayudarte a levantar el peso.')
add_bullet(doc, 'El movimiento debe ser controlado tanto al subir como al bajar.')

hr(doc)

# ── 6.3 PRESS DE HOMBRO ─────────────────────
add_heading(doc, '6.3 Press de Hombro', level=2)
add_screenshot_placeholder(doc, 'chip "Press de Hombro" activo con barra verde "¡Extensión completa!"')

add_heading(doc, 'Configuración de la cámara', level=3)
add_bullet(doc, 'Posición recomendada: Cámara de frente o ligeramente de costado, a la altura del pecho o la cintura.')
add_bullet(doc, 'Distancia: Al menos 1.5 metros para que hombros, codos y muñecas sean visibles por encima de la cabeza.')
add_bullet(doc, 'Articulaciones que monitorea: Hombros, codos y muñecas de ambos brazos.')
add_note(doc,
    'El movimiento de press es opuesto al curl: el esfuerzo sube las pesas sobre la cabeza, '
    'extendiendo el codo. La app detecta este patrón automáticamente.',
    kind='note',
)

add_heading(doc, 'Cómo ejecutar el press de hombro', level=3)
add_screenshot_placeholder(doc, 'posición inicial con pesas a nivel de hombros y barra idle "Listo — empuja hacia arriba"')
for i, txt in enumerate([
    'Sostén el peso (mancuernas, barra o banda) a la altura de los hombros con los codos doblados.',
    'La barra inferior mostrará: "Listo — empuja hacia arriba".',
    'Empuja el peso hacia arriba extendiendo los brazos sobre la cabeza.',
], 1):
    add_numbered(doc, txt, i)
add_screenshot_placeholder(doc, 'pesas sobre la cabeza con barra verde "¡Extensión completa!"')
for i, txt in enumerate([
    'Cuando el ángulo del codo supera los 145°, la barra muestra "¡Extensión completa!" (verde).',
    'Si la barra muestra "Extiende un poco más" (amarillo), el ángulo del codo está entre 100° y 145°.',
    'Baja el peso controladamente hasta la posición inicial. La repetición se contabiliza en ese momento.',
], 4):
    add_numbered(doc, txt, i)
add_screenshot_placeholder(doc, 'posición de bajada excesiva con barra ROJA "No bajes tanto — cuida los hombros"')
add_note(doc,
    'Alerta de seguridad: Si el ángulo del codo baja de 80° (codos más bajos que los hombros con carga), '
    'la barra muestra en ROJO "No bajes tanto — cuida los hombros". Esta posición puede comprimir el tendón '
    'supraespinoso y causar lesión. Sube el peso inmediatamente.',
    kind='alert',
)

add_heading(doc, 'Tabla de retroalimentación — Press de Hombro', level=3)
make_table(doc,
    headers=['Estado', 'Color', 'Mensaje', 'Significado'],
    rows=[
        ['Pesas a nivel de hombros',          'Blanco',   '"Listo — empuja hacia arriba"',              'Posición inicial correcta'],
        ['Overhead, extensión incompleta',    'Amarillo', '"Extiende un poco más"',                     'Ángulo del codo entre 100° y 145°'],
        ['Overhead, extensión completa',      'Verde',    '"¡Extensión completa!"',                     'Ángulo del codo ≥ 145°'],
        ['Bajada excesiva — riesgo',          'Rojo',     '"No bajes tanto — cuida los hombros"',       'Ángulo del codo < 80°'],
        ['Brazo fuera de encuadre',           'Blanco',   '"Asegúrate de que tu brazo sea visible"',    'Articulaciones no detectadas'],
    ],
    col_widths=[4.5, 2.5, 5, 4],
)

add_heading(doc, 'Consejos de posición', level=3)
add_bullet(doc, 'La espalda baja debe estar neutral (sin arquearse hacia atrás).')
add_bullet(doc, 'Los codos deben apuntar hacia adelante, no hacia los lados, en la posición inicial.')
add_bullet(doc, 'Evitá hiperextender los codos al llegar arriba (no los bloquees completamente).')

# ══════════════════════════════════════════════
# 7. RETROALIMENTACIÓN
# ══════════════════════════════════════════════
add_heading(doc, '7. Sistema de retroalimentación', level=1)
add_body(doc, 'La aplicación te da retroalimentación por dos canales simultáneos: visual y auditivo.')

add_heading(doc, '7.1 Retroalimentación visual', level=2)
add_screenshot_placeholder(doc, 'barra de retroalimentación en tres estados lado a lado: verde, amarillo y rojo')
add_body(doc, 'La barra inferior cambia de color según la calidad de tu ejecución:')
make_table(doc,
    headers=['Color del borde', 'Nivel', 'Significado general'],
    rows=[
        ['**Verde**',                   'Buena forma',         'Ejecución técnicamente correcta'],
        ['**Amarillo**',                'Forma mejorable',     'Podés ajustar para mejor resultado'],
        ['**Rojo**',                    'Posición de riesgo',  'Detente o ajusta inmediatamente'],
        ['**Blanco semitransparente**', 'Inactivo',            'Esperando inicio de movimiento o cuerpo no detectado'],
    ],
    col_widths=[4.5, 4, 7.5],
)
add_note(doc,
    'El fondo de la barra usa un efecto de desenfoque (cristal esmerilado) para ser legible '
    'sobre cualquier fondo de video.',
    kind='note',
)

add_heading(doc, '7.2 Retroalimentación de voz', level=2)
add_body(doc,
    'La aplicación habla en español para que puedas mantener la vista en el entorno sin necesidad '
    'de leer la pantalla. La voz se activa en los siguientes momentos:'
)
make_table(doc,
    headers=['Evento', 'Lo que dice'],
    rows=[
        ['Al completar 1 repetición',                               '"Una"'],
        ['Al completar 5 repeticiones',                             '"Cinco. ¡Sigue así!"'],
        ['Al completar 10, 20, 30... repeticiones',                 '"Diez. ¡Excelente ritmo!"'],
        ['Cualquier otro número de reps',                           'El número (ej. "Dos", "Tres")'],
        ['Fondo de sentadilla con profundidad óptima',              '"¡Excelente profundidad!"'],
        ['Fondo de sentadilla con profundidad insuficiente',        '"Baja un poco más"'],
        ['Rep de curl con buena contracción',                       '"[número]. ¡Excelente contracción!"'],
        ['Rep de curl con contracción incompleta',                  '"[número]. Sube un poco más"'],
        ['Rep de press con buena extensión',                        '"[número]. ¡Extensión completa!"'],
        ['Rep de press con extensión incompleta',                   '"[número]. Extiende un poco más"'],
    ],
    col_widths=[8, 8],
)
add_note(doc, 'Para mejor experiencia de voz, usa auriculares o sube el volumen del dispositivo.', kind='warning')
add_note(doc,
    'iOS Safari: La voz funciona correctamente después de haber tocado el botón "Comenzar a entrenar" '
    'en el onboarding. Si la voz no suena, verifica que el modo silencioso del iPhone no esté activado.',
    kind='warning',
)

# ══════════════════════════════════════════════
# 8. CAMBIAR EJERCICIO
# ══════════════════════════════════════════════
add_heading(doc, '8. Cambiar de ejercicio', level=1)
add_screenshot_placeholder(doc, 'secuencia de tap en chip "Curl de Bíceps" con el contador reseteado a 0')
add_body(doc, 'Para cambiar de ejercicio en cualquier momento:')
for i, txt in enumerate([
    'Toca el chip del ejercicio que deseas en la barra inferior.',
    'El chip seleccionado se resalta.',
    'El contador de repeticiones se resetea a 0 automáticamente.',
    'El esqueleto y la retroalimentación se adaptan al nuevo ejercicio de inmediato.',
], 1):
    add_numbered(doc, txt, i)
add_note(doc,
    'Al cambiar de ejercicio se pierde el conteo acumulado del ejercicio anterior. '
    'Si necesitás anotar tus repeticiones, hacelo antes de cambiar.',
    kind='alert',
)

# ══════════════════════════════════════════════
# 9. CAMBIAR CÁMARA
# ══════════════════════════════════════════════
add_heading(doc, '9. Cambiar de cámara durante el entrenamiento', level=1)
add_screenshot_placeholder(doc, 'botón de cambio de cámara tocado con mensaje "Cambiando cámara..." visible')
for i, txt in enumerate([
    'Toca el botón circular con el ícono de cámara en la esquina inferior derecha.',
    'Aparece brevemente el mensaje "Cambiando cámara..." mientras se libera la cámara anterior.',
    'En 1–2 segundos, la imagen cambia a la cámara seleccionada.',
    'El conteo de repeticiones no se resetea al cambiar de cámara.',
], 1):
    add_numbered(doc, txt, i)
add_note(doc,
    'El breve tiempo de espera al cambiar de cámara es normal. Es más notable en dispositivos '
    'de gama baja y en modo PWA instalada, donde el hardware necesita liberar el sensor.',
    kind='note',
)

# ══════════════════════════════════════════════
# 10. SOLUCIÓN DE PROBLEMAS
# ══════════════════════════════════════════════
add_heading(doc, '10. Solución de problemas', level=1)

add_heading(doc, 'La cámara no inicia o muestra un error', level=2)
add_screenshot_placeholder(doc, 'pantalla de error con mensaje "Error al iniciar la cámara: ..."')
make_table(doc,
    headers=['Problema', 'Causa probable', 'Solución'],
    rows=[
        ['"Contexto no seguro: abrí la app con HTTPS"', 'Acceso por HTTP',            'Usa la URL correcta que empieza con https://'],
        ['"Error al iniciar la cámara" genérico',        'Permisos denegados',         'Ver procedimiento de permisos en la siguiente sección'],
        ['Pantalla negra sin error',                     'Hardware ocupado por otra app', 'Cierra otras apps que usen la cámara y recarga la página'],
        ['Cámara inicia pero sin esqueleto',             'Iluminación insuficiente',   'Mejora la iluminación del ambiente'],
    ],
    col_widths=[5.5, 4, 6.5],
)

add_heading(doc, 'Habilitar permisos de cámara si fueron denegados', level=2)

add_body_mixed(doc, [('En Android (Chrome):', True, False)])
add_screenshot_placeholder(doc, 'configuración de permisos de sitio en Chrome Android (ícono de candado → Permisos → Cámara)')
for i, txt in enumerate([
    'Toca el ícono de candado o ⓘ en la barra de dirección del navegador.',
    'Toca "Permisos".',
    'Activa la opción "Cámara".',
    'Recarga la página.',
], 1):
    add_numbered(doc, txt, i)

add_body_mixed(doc, [('En iOS (Safari):', True, False)])
add_screenshot_placeholder(doc, 'configuración de Safari en Ajustes de iOS (Ajustes → Safari → Cámara → Permitir)')
for i, txt in enumerate([
    'Ve a Ajustes del iPhone.',
    'Desplázate hacia abajo y toca "Safari".',
    'Toca "Cámara" y selecciona "Permitir".',
    'Regresa a Safari y recarga la página de la aplicación.',
], 1):
    add_numbered(doc, txt, i)

add_heading(doc, 'El esqueleto no aparece o aparece distorsionado', level=2)
add_screenshot_placeholder(doc, 'barra inferior con mensaje "Asegúrate de que tu cuerpo completo sea visible"')
make_table(doc,
    headers=['Problema', 'Solución'],
    rows=[
        ['Cuerpo fuera del encuadre',                         'Aléjate de la cámara hasta que todo el cuerpo sea visible'],
        ['Iluminación insuficiente',                           'Enciende más luces o acércate a una fuente de luz'],
        ['Ropa muy holgada que oculta la silueta',            'Usa ropa ajustada o de colores que contrasten con el fondo'],
        ['Fondo oscuro o con muchos elementos visuales',      'Busca una pared lisa de un solo color como fondo'],
        ['Varias personas en el encuadre',                    'Asegúrate de ser la única persona visible en la cámara'],
    ],
    col_widths=[7, 9],
)

add_heading(doc, 'La voz no funciona', level=2)
make_table(doc,
    headers=['Problema', 'Solución'],
    rows=[
        ['Voz en silencio en iPhone',             'Verifica que el switch físico de silencio del iPhone esté desactivado'],
        ['La voz habla en otro idioma',           'Ve a Ajustes → Accesibilidad → Contenido hablado → Voces y descarga la voz de español'],
        ['La voz nunca inició',                   'Asegúrate de haber tocado el botón "Comenzar a entrenar" en el onboarding (requerido por iOS)'],
        ['No se escucha aunque hay volumen',      'Sube el volumen multimedia del dispositivo (distinto al volumen de llamadas)'],
    ],
    col_widths=[7, 9],
)

add_heading(doc, 'El conteo de repeticiones no es preciso', level=2)
make_table(doc,
    headers=['Síntoma', 'Posible causa', 'Solución'],
    rows=[
        ['Cuenta de más',           'Movimientos bruscos o cámara inestable',              'Asegura bien el soporte de la cámara'],
        ['No cuenta algunas reps',  'Rango de movimiento incompleto',                      'Completa el ciclo completo del ejercicio (bajada total + subida total)'],
        ['Cuenta menos (curl)',     'Vista lateral sin el brazo correcto visible',          'Reposiciona la cámara o usa vista frontal'],
    ],
    col_widths=[4.5, 5.5, 6],
)

# ══════════════════════════════════════════════
# 11. FAQ
# ══════════════════════════════════════════════
add_heading(doc, '11. Preguntas frecuentes', level=1)

faqs = [
    ('¿Funciona sin conexión a internet?',
     'Sí. Una vez que hayas abierto la app por primera vez y descargado el modelo de IA (~8 MB), '
     'la aplicación funciona completamente sin conexión. La PWA instalada cachea todos los recursos necesarios.'),
    ('¿El video de mi cámara se envía a algún servidor?',
     'No. Todo el procesamiento de video ocurre en tu dispositivo. '
     'La imagen de tu cámara nunca sale de tu celular.'),
    ('¿Por qué la primera carga tarda más que las siguientes?',
     'En la primera apertura, el navegador descarga el modelo de IA de MediaPipe (~8 MB) y lo guarda '
     'en el caché del dispositivo. Las aperturas siguientes son casi instantáneas porque el modelo '
     'ya está almacenado localmente.'),
    ('¿Funciona en tablet o computadora?',
     'La aplicación está diseñada para celulares. Puede funcionar en tablet con resultados similares. '
     'En computadoras de escritorio puede abrirse en el navegador, pero la experiencia no está '
     'optimizada para ese formato y el acceso a la cámara trasera no aplica.'),
    ('¿El conteo se guarda entre sesiones?',
     'La versión actual no guarda el historial de sesiones. '
     'Cada vez que abres la app, el contador comienza en cero.'),
    ('¿Por qué la app me pide permiso de cámara en cada navegador?',
     'Los permisos de cámara son por dominio y por navegador. Si instalás la PWA desde Chrome, '
     'los permisos otorgados en Chrome se aplican también a la PWA instalada. Si abrís la misma URL '
     'en otro navegador, deberás otorgar los permisos nuevamente.'),
    ('¿Puedo usar la app en modo oscuro?',
     'La interfaz de la aplicación no cambia con el modo oscuro del sistema operativo. '
     'La pantalla de la cámara tiene contraste alto para ser visible en exteriores o ambientes bien iluminados.'),
    ('¿Qué hago si la aplicación se congela?',
     '1) Cierra la pestaña o la app completamente. '
     '2) Vuelve a abrirla. '
     '3) Si el problema persiste, intenta limpiar el caché del navegador o reinstalar la PWA.'),
]
for question, answer in faqs:
    add_body_mixed(doc, [(question, True, False)])
    add_body(doc, answer, indent=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

hr(doc)

# Nota de pie de documento
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(
    'Documento elaborado para la entrega del curso Inteligencia Artificial — IA26.\n'
    'Universidad Mariano Gálvez de Guatemala — Mayo 2026.'
)
set_font(r, size=10, italic=True, color=(120, 120, 120))

# ── GUARDAR ──────────────────────────────────
out = os.path.abspath(OUTPUT_PATH)
doc.save(out)
print(f'Documento guardado: {out}')
